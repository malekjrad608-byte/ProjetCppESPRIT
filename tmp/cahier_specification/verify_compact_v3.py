from pathlib import Path
from collections import Counter
from hashlib import sha256
import re
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

root = Path(r'C:\Users\malek\Documents\gestion_service')
new = Document(root/'output/documents/ANIWORLD_devCore_Cahier_de_specification_v3_compact.docx')
old = Document(root/'output/documents/ANIWORLD_devCore_Cahier_de_specification_v2.docx')
pdf = PdfReader(root/'output/pdf/ANIWORLD_devCore_Cahier_de_specification_v3_compact.pdf')
text = '\n'.join(p.extract_text() for p in pdf.pages)
assert len(pdf.pages) == 12
base_ids = {f'{prefix}-{i:02}' for prefix in ['SER','EMP','STK','LOC','DM'] for i in range(1,6)}
innovations = {f'{prefix}-I{i}' for prefix in ['SER','EMP','STK','LOC','DM'] for i in (1,2)}
quality = {f'NF-{i:02}' for i in range(1,12)}
for identifier in base_ids | innovations | quality:
    assert identifier in text, identifier
def table_text(t):
    return [[c.text for c in r.cells] for r in t.rows]
for source_index in range(4,9):
    target = table_text(old.tables[source_index])
    assert any(table_text(t) == target for t in new.tables), source_index
headings = [p.text for p in new.paragraphs if p.style.name.startswith('Heading')]
for prefix in ['1 Contexte', '1.1 ', '1.2 ', '2 La solution', '2.1 ', '2.2 ', '2.3 ', '3 Conception', '3.1 ', '3.2 ', '3.3 ']:
    assert any(h.startswith(prefix) for h in headings), prefix
assert 'Sami Ben Salem' in pdf.pages[2].extract_text()
assert 'fictifs' in pdf.pages[2].extract_text()
assert 'Les Compagnons' in text
assert len(new.inline_shapes) == 10
hashes = Counter()
for shape in new.inline_shapes:
    rid = shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
    hashes[sha256(new.part.related_parts[rid].blob).hexdigest()] += 1
assets = Path(r'C:\Users\malek\Downloads')
for name, count in [('logo-removebg-preview.png',1),('new logo.png',1),('new name.png',2)]:
    assert hashes[sha256((assets/name).read_bytes()).hexdigest()] == count, name
for i in range(6):
    gui = root/f'tmp/cahier_specification/gui_aniworld_v2/gui_{i}.png'
    assert hashes[sha256(gui.read_bytes()).hexdigest()] == 1
print('PASS: 12 pages; 25 base requirements unchanged; 10 innovations; 11 quality requirements; all mandatory sections; fictional client; group logo once, correct application logos and six interfaces.')
