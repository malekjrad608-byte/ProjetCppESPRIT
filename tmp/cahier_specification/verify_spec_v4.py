from pathlib import Path
from hashlib import sha256
from collections import Counter
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

root=Path(r'C:\Users\malek\Documents\gestion_service')
doc=Document(root/'output/documents/ANIWORLD_DevCore_Cahier_de_specification_v4.docx')
source=Document(root/'output/documents/ANIWORLD_devCore_Cahier_de_specification_v3_compact.docx')
reader=PdfReader(root/'output/pdf/ANIWORLD_DevCore_Cahier_de_specification_v4.pdf')
text='\n'.join(p.extract_text() for p in reader.pages)
assert len(reader.pages)==14
for module in ('SER','EMP','STK','LOC','DM','ADOP'):
    for i in range(1,6):assert f'{module}-{i:02}' in text
    for i in (1,2):assert f'{module}-I{i}' in text
for i in range(1,12):assert f'NF-{i:02}' in text
def values(t):return [[c.text for c in r.cells] for r in t.rows]
for table in source.tables[3:8]:assert any(values(t)==values(table) for t in doc.tables)
for phrase in ('chat orange','patte verte','aux cinq modules','les cinq gestions','les dix métiers','devCore'):
    assert phrase not in text,phrase
headings=[p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
for prefix in ('1 Contexte','1.1 ','1.2 ','2 La solution','2.1 ','2.2 ','2.3 ','3 Conception','3.1 ','3.2 ','3.3 '):
    assert any(p.startswith(prefix) for p in headings),prefix
assert 'Sami Ben Salem' in text and 'fictifs' in text
assert len(doc.inline_shapes)==10
hashes=Counter()
for shape in doc.inline_shapes:
    rid=shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
    hashes[sha256(doc.part.related_parts[rid].blob).digest()]+=1
downloads=Path(r'C:\Users\malek\Downloads')
for name,count in (
    ('Plan de travail 1.png',2),
    ('WhatsApp Image 2026-09-30 at 11.27.16 PM.jpeg',1),
    ('ChatGPT Image Oct 1, 2026, 11_16_40 AM.png',1),
    ('WhatsApp Image 2026-09-30 at 9.04.54 PM.jpeg',1),
    ('WhatsApp Image 2026-09-30 at 9.11.47 PM.jpeg',1),
    ('5604730f-5391-4456-8296-8146c9a70582.jpg',1)):
    assert hashes[sha256((downloads/name).read_bytes()).digest()]==count,name
for i in (0,5,6):
    path=root/f'tmp/cahier_specification/gui_aniworld_v4/gui_{i}.png'
    assert hashes[sha256(path.read_bytes()).digest()]==1
assert sum(hashes.values())==10
assert 'Cycle : En étude, Validée, Refusée ou Annulée.' in text
print('PASS: 14 pages; all required plan sections; correct ANIWORLD and DevCore logos; four supplied screenshots unchanged; new adoption interface; 30 base requirements, 12 innovations and 11 quality requirements; existing five base tables preserved.')
