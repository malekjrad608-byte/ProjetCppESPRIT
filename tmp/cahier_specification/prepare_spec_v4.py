from pathlib import Path
work=Path(__file__).parent
s=(work/'build_compact_v3.py').read_text(encoding='utf-8')
def edit(a,b):
 global s
 assert a in s,a[:100]
 s=s.replace(a,b)
edit("SOURCE=ROOT/'output/documents/ANIWORLD_devCore_Cahier_de_specification_v2.docx'","SOURCE=ROOT/'output/documents/ANIWORLD_devCore_Cahier_de_specification_v3_compact.docx'")
edit("DEST=ROOT/'output/documents/ANIWORLD_devCore_Cahier_de_specification_v3_compact.docx'","DEST=ROOT/'output/documents/ANIWORLD_DevCore_Cahier_de_specification_v4.docx'")
edit("t.text=t.text.replace('Version 1.2','Version 1.3')","t.text=t.text.replace('Version 1.3','Version 1.4').replace('devCore','DevCore').replace('27 septembre 2026','1 octobre 2026')")
edit("d.core_properties.author='devCore'","d.core_properties.author='DevCore'")
edit('Groupe devCore','Groupe DevCore')
edit("table(oldtables[3][0],oldtables[3][1:],[2.2,1.65,4.05,2.65,7.05],9.3)","""overview=oldtables[2][1:]+[['Gestion des Adoptions','Adoption','ID, animal, candidat, foyer, date, statut, référent ; validation et suivi.','Référent adoption ; responsable.','CRUD ; tri date/statut ; recherche ; export PDF/CSV ; statistiques par statut.\\nADOP-I1 : compatibilité avec le foyer.\\nADOP-I2 : suivi après adoption.']]
table(oldtables[2][0],overview,[2.2,1.65,4.05,2.65,7.05],9.2)""")
for a,b in [("('2.1.1 Gestion des Services',4,oldtext[71]","('2.1.1 Gestion des Services',3,oldtext[62]"),("('2.1.2 Gestion des Employés',5,oldtext[82]","('2.1.2 Gestion des Employés',4,oldtext[68]"),("('2.1.3 Gestion du Stock',6,oldtext[93]","('2.1.3 Gestion du Stock',5,oldtext[75]"),("('2.1.4 Gestion des Locaux',7,oldtext[104]","('2.1.4 Gestion des Locaux',6,oldtext[81]"),("('2.1.5 Gestion des dossiers médicaux',8,oldtext[115]","('2.1.5 Gestion des dossiers médicaux',7,oldtext[88]")]:edit(a,b)
edit('aux cinq modules','aux six modules')
edit('les cinq gestions','les six gestions')
edit('les conventions des cinq modules','les conventions des six modules')
edit('les dix métiers innovants','les douze métiers innovants')
edit('Le groupe devCore propose','Le groupe DevCore propose')
edit('M. Ben Salem coordonne huit collaborateurs et six espaces.','M. Ben Salem coordonne une équipe pluridisciplinaire et plusieurs espaces ; le centre accompagne aussi des adoptions.')
edit('Les disponibilités des salles sont souvent communiquées oralement.','Les disponibilités des salles sont souvent communiquées oralement ; les candidatures à l’adoption sont suivies dans des échanges séparés.')
edit('Il manque aussi d’indicateurs fiables pour préparer sa journée et répartir le travail.','Il manque aussi d’indicateurs fiables pour préparer sa journée, répartir le travail et suivre les demandes d’adoption après le départ d’un animal.')
edit("],[4.4,8.0,5.2],9.5)","""    ['Candidatures à l’adoption dispersées et suivi après départ irrégulier.','Dossier de candidature, comparaison animal–foyer, décision validée et suivi à J+7 / J+30.','Organiser les placements et garder une trace des échanges avec les adoptants.'],
],[4.4,8.0,5.2],9.3)""")
edit('Le vétérinaire consulte et complète le contenu clinique. L’administrateur','Le vétérinaire consulte et complète le contenu clinique. Le référent adoption instruit les candidatures ; le responsable valide le placement. L’administrateur')
adoption="""    ('2.1.6 Gestion des Adoptions',None,
     'Adoption : idAdoption, idAnimal, idCandidat, dateDemande, statut, idRéférent, dateDécision, datePlacement. Candidat : identité, contact, logement, expérience et composition du foyer. Données associées : pièces, rencontre, décision et suivis.',
     'ADOP-I1 Compatibilité avec le foyer. Comparer les besoins renseignés de l’animal aux caractéristiques du candidat : logement, activité, enfants et autres animaux. Afficher les critères concordants, manquants et leur proportion ; le responsable prend la décision.',
     'ADOP-I2 Suivi après adoption. À la date de placement, créer deux actions à J+7 et J+30. Enregistrer contact, réponse et résultat ; signaler les actions en retard au référent, qui décide de la suite à donner.',
     'Règles. Cycle : En étude, Validée, Refusée ou Annulée. Une seule adoption validée et active par animal ; décision et placement datés. Vérifier pièces, rencontre et accord du responsable avant validation. Le référent voit une synthèse autorisée sans accès clinique.',
     'Recette. Pour 3 critères concordants sur 4 évaluables, afficher 75 % et le critère restant. Bloquer une seconde adoption active ; pour un placement le 05/10, créer des suivis les 12/10 et 04/11.'),
"""
edit("\n]\nfor i,(heading,ti,attrs,innovation1,innovation2,rules,test) in enumerate(module_data):",'\n'+adoption+"]\nfor i,(heading,ti,attrs,innovation1,innovation2,rules,test) in enumerate(module_data):")
edit('if i in [0,2,4]:page()','if i in [0,2,4,5]:page()')
edit('table(oldtables[ti][0],oldtables[ti][1:],[3.3,14.3],9.5)',"""if ti is None:
        table(['Exigence','Fonctionnalité attendue'],[
            ['ADOP-01 CRUD','Créer une demande ; consulter le dossier ; modifier avant décision ; supprimer un brouillon ou archiver une demande traitée.'],
            ['ADOP-02 Trier','Trier par date de demande, animal, candidat ou statut, dans les deux sens.'],
            ['ADOP-03 Rechercher','Rechercher par dossier, animal ou candidat ; filtrer par espèce, référent, période et statut.'],
            ['ADOP-04 Exporter','Exporter un dossier autorisé en PDF et une liste filtrée en CSV ; limiter les données personnelles au besoin du destinataire.'],
            ['ADOP-05 Statistiques','Présenter les demandes par statut, les placements par période et la part de suivis réalisés à échéance.'],
        ],[3.3,14.3],9.5)
    else:table(oldtables[ti][0],oldtables[ti][1:],[3.3,14.3],9.5)""")
edit('if i==4:\n        h(\'2.1.6 Échanges entre les modules\',3)','if i==5:\n        h(\'2.1.7 Échanges entre les modules\',3)')
edit("['Locaux ↔ Employés','Identifier l’agent ayant effectué ou validé le nettoyage et la maintenance.'],","""['Locaux ↔ Employés','Identifier l’agent ayant effectué ou validé le nettoyage et la maintenance.'],
            ['Adoptions ↔ DM / Employés','Réutiliser l’identité de l’animal, une synthèse autorisée et le référent ; limiter l’accès clinique. Une adoption validée change l’état de placement.'],""")
edit('Le tableau de bord actualise l’activité ; le journal conserve les actions.','Le référent suit les candidatures à l’adoption puis les contacts après placement. Le tableau de bord actualise l’activité ; le journal conserve les actions.')
edit('rappels et poids (DM).','rappels et poids (DM), compatibilité et suivi après placement (Adoptions).')
edit('Prévoir un cas nominal et un cas de refus ou de limite pour chaque CRUD et métier innovant.','Prévoir un cas nominal et un cas de refus ou de limite pour chaque CRUD et métier innovant, y compris une adoption incomplète ou déjà validée.')
edit("for name,width in [('new logo.png',1.9),('new name.png',4.7)]:","for name,width in [('Plan de travail 1.png',1.9)]:")
edit('ANIWORLD utilise le chat orange dans un médaillon bleu et le logotype fourni. La patte verte appartient au groupe devCore et figure uniquement sur la couverture. Conserver les couleurs, la transparence et les proportions ; réserver un espace libre sur fond clair.','ANIWORLD utilise le logo chien–chat entouré de mains, de feuilles et d’une croix orange, accompagné du nom ANIWORLD. Le logo bleu et blanc fourni représente le groupe DevCore sur la couverture. Conserver les proportions et les couleurs de chaque identité, avec un espace libre autour.')
edit('Rappel du chat et encarts de suivi.','Croix du logo et encarts de suivi.')
edit('Captures de maquettes Qt Widgets avec données fictives. Elles illustrent les fonctionnalités décrites dans ce cahier.','Écrans avec données de démonstration. Services, Employés, Stock et Locaux reprennent les images fournies ; Tableau de bord, Dossiers médicaux et Adoptions sont des maquettes Qt.')
edit('Le planning relie chaque prestation à un horaire et un local. La fiche présente l’animal, l’employé et le tarif ; un créneau proposé est recontrôlé avant confirmation.','La vue en cartes présente les prestations, dates, durées et statuts. Les filtres ouvrent la fiche sélectionnée avec son responsable, son tarif et les actions de consultation, modification et démarrage.')
edit('La recherche ouvre la fiche, les compétences et les disponibilités de l’employé. La charge hebdomadaire et la suggestion de réaffectation aident à répartir le travail.','Les cartes présentent les employés, leurs disponibilités, compétences et heures planifiées. Le panneau détaille la fiche choisie ; recherche, filtres, modification et statistiques facilitent la répartition du travail.')
edit('Les cartes montrent les capacités et les états. La remise en service exige le nettoyage validé ; la maintenance réserve sa période d’indisponibilité.','Les cartes présentent les types, capacités et états des locaux. Le panneau du local sélectionné regroupe affectation, libération, modification et suppression ; la vue d’occupation complète la recherche et les filtres.')
edit(" ('3.3.6 Gestion des dossiers médicaux','Le vétérinaire retrouve l’animal, les consultations, les pesées et les échéances. Les rappels restent à valider ; le signal de poids demande une vérification clinique.'),",""" ('3.3.6 Gestion des dossiers médicaux','Le vétérinaire retrouve l’animal, les consultations, les pesées et les échéances. Les rappels restent à valider ; le signal de poids demande une vérification clinique.'),
 ('3.3.7 Gestion des Adoptions','Les cartes relient un animal à sa candidature. Le panneau affiche les critères de compatibilité, les étapes de validation et le suivi à J+7 / J+30 ; le responsable confirme le placement.'),""")
edit("path=WORK/'gui_aniworld_v2'/f'gui_{i}.png'","path=gui_paths[i]")
edit("    shape=ip.add_run().add_picture(str(path),width=Cm(16.0))","""    shape=ip.add_run().add_picture(str(path),width=Cm(16.0))
    if i==4:crop_to(shape,path,(0,24,1600,950),16.0)
    if i==3:crop_to(shape,path,(0,20,1600,857),16.0)
    if i==2:crop_to(shape,path,(0,0,1600,835),16.0)""")
edit("for i,(title,explanation) in enumerate(gui):","""gui_paths=[
    WORK/'gui_aniworld_v4/gui_0.png',
    Path(r'C:\\Users\\malek\\Downloads\\ChatGPT Image Oct 1, 2026, 11_16_40 AM.png'),
    Path(r'C:\\Users\\malek\\Downloads\\WhatsApp Image 2026-09-30 at 9.04.54 PM.jpeg'),
    Path(r'C:\\Users\\malek\\Downloads\\5604730f-5391-4456-8296-8146c9a70582.jpg'),
    Path(r'C:\\Users\\malek\\Downloads\\WhatsApp Image 2026-09-30 at 9.11.47 PM.jpeg'),
    WORK/'gui_aniworld_v4/gui_5.png',
    WORK/'gui_aniworld_v4/gui_6.png',
]
for i,(title,explanation) in enumerate(gui):""")
edit("d.save(DEST)","""h('Parcours illustré de l’adoption',4)
table(['Étape','Action et résultat attendu'],[
    ['1 · Demande','L’accueil enregistre le candidat et l’animal ; le dossier passe à En étude.'],
    ['2 · Vérification','Le référent vérifie les pièces, la rencontre et les critères de compatibilité.'],
    ['3 · Décision','Le responsable valide ou refuse avec motif ; une adoption validée conserve son historique.'],
    ['4 · Suivi','La date de placement déclenche les contacts à J+7 et J+30 ; chaque réponse est tracée.'],
],[3.5,14.1],9.5)
p('Cette maquette illustre un dossier encore en étude : la validation est inactive tant que la rencontre et l’accord du responsable ne sont pas enregistrés.',size=9.5,after=0)
d.save(DEST)""")
# Extend crop support to opaque supplied artwork and to desktop screenshot frames.
edit('from PIL import Image','from PIL import Image, ImageChops')
edit('im=Image.open(path);w,h=im.size;l,t,r,b=im.getbbox()',"""im=Image.open(path);w,h=im.size
    if name_is_app(path):l,t,r,b=ImageChops.difference(im.convert('RGB'),Image.new('RGB',im.size,'white')).getbbox()
    else:l,t,r,b=im.getbbox()""")
marker='# Page 3: one named fictional client, his complete problem and the expected benefit.'
cover=r'''
def name_is_app(path):return Path(path).name=='Plan de travail 1.png'
def crop_to(shape,path,box,width):
    w,h=Image.open(path).size;l,t,r,b=box
    shape.width=Cm(width);shape.height=Cm(width*(b-t)/(r-l))
    ext=shape._inline.xpath('.//a:xfrm/a:ext')[0];ext.set('cx',str(shape.width));ext.set('cy',str(shape.height))
    fill=shape._inline.xpath('.//a:blip')[0].getparent()
    for x in list(fill.findall(qn('a:srcRect'))):fill.remove(x)
    crop=OxmlElement('a:srcRect')
    for k,v in [('l',l/w),('t',t/h),('r',(w-r)/w),('b',(h-b)/h)]:crop.set(k,str(round(v*100000)))
    fill.insert(1,crop)
coverparas=d.paragraphs
group=coverparas[3];group.clear();group.alignment=WD_ALIGN_PARAGRAPH.CENTER
group.add_run().add_picture(r'C:\Users\malek\Downloads\WhatsApp Image 2026-09-30 at 11.27.16 PM.jpeg',width=Cm(5.0))
for run in coverparas[4].runs:run.text=run.text.replace('devCore','DevCore')
brand=coverparas[5];brand.clear();brand.alignment=WD_ALIGN_PARAGRAPH.CENTER
shape=brand.add_run().add_picture(r'C:\Users\malek\Downloads\Plan de travail 1.png');crop_shape(shape,Path(r'C:\Users\malek\Downloads\Plan de travail 1.png'),2.5)
brand.paragraph_format.space_after=Pt(2);brand.paragraph_format.keep_with_next=True
from docx.text.paragraph import Paragraph
name=OxmlElement('w:p');brand._p.addnext(name);np=Paragraph(name,brand._parent);np.alignment=WD_ALIGN_PARAGRAPH.CENTER
nr=np.add_run('ANIWORLD');nr.bold=True;nr.font.size=Pt(22);nr.font.color.rgb=RGBColor.from_string('263B42');np.paragraph_format.space_after=Pt(8)
coverparas[8].text='Services  •  Employés  •  Stock\nLocaux  •  Dossiers médicaux  •  Adoptions'
coverparas[8].alignment=WD_ALIGN_PARAGRAPH.CENTER
for run in coverparas[8].runs:run.font.size=Pt(10.5)
for cp in d.paragraphs:
    if cp.text.startswith('Les identifiants SER, EMP, STK, LOC et DM'):
        cp.text='Les identifiants SER, EMP, STK, LOC, DM et ADOP relient les exigences aux six modules. Les critères de recette décrivent les résultats attendus lors de la démonstration.'
'''
edit(marker,cover+'\n'+marker)
(work/'build_spec_v4.py').write_text(s,encoding='utf-8')
print('Prepared compact six-module specification with corrected identities and supplied screenshots.')
