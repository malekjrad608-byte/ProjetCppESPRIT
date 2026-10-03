from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\malek\Documents\gestion_service')
WORK=ROOT/'tmp'/'cahier_specification'
SOURCE=ROOT/'output'/'documents'/'Pattia_Cahier_de_specification.docx'
DEST=ROOT/'output'/'documents'/'ANIWORLD_Cahier_de_specification.docx'
NEWLOGO=Path(r'C:\Users\malek\Downloads\new name.png')
doc=Document(SOURCE)
for part in doc.part.package.parts:
    if hasattr(part,'element'):
        for t in part.element.iter(qn('w:t')):
            if t.text:t.text=t.text.replace('Pattia','ANIWORLD').replace('PATTIA','ANIWORLD')
doc.core_properties.title='Cahier de spécification ANIWORLD'
doc.core_properties.author='Équipe projet ANIWORLD'
doc.core_properties.keywords=doc.core_properties.keywords.replace('Pattia','ANIWORLD')

def paragraphs():
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:yield from cell.paragraphs

def replace(old,new):
    matches=0
    for p in paragraphs():
        if p.text==old:
            if p.runs:
                p.runs[0].text=new
                for run in p.runs[1:]:run.text=''
            else:p.add_run(new)
            matches+=1
    if not matches:raise ValueError('Missing paragraph: '+old)

replace('L’identité visuelle de ANIWORLD associe une patte animale, un vert sauge et une interface claire. Elle accompagne une application de travail quotidien : les actions et les états doivent rester faciles à distinguer.',
        'L’identité visuelle d’ANIWORLD associe le logotype fourni, une mascotte féline et une interface aux tons verts. Elle accompagne une application de travail quotidien : les actions et les états doivent rester faciles à distinguer.')
replace('Figure 1  •  Logo de ANIWORLD fourni pour le projet','Figure 1  •  Logotype ANIWORLD fourni pour le projet')
replace('Le nom ANIWORLD évoque la patte représentée dans le logo. Il est court et permet d’identifier l’application sans la limiter à une seule prestation. Le symbole renvoie aux animaux pris en charge ; le cercle vert donne une forme simple et reconnaissable à l’identité du centre.',
        'Le nom ANIWORLD évoque l’univers des animaux. Le logotype associe des lettres sombres, une patte et un chat orange dans un médaillon bleu. La mascotte reprend ce chat dans les bandeaux de l’interface et crée un repère commun aux cinq modules.')
replace('Patte anthracite','Logotype ANIWORLD')
replace('Symbole principal lié à l’activité du centre. Le contraste assure sa lisibilité sur le cercle vert.',
        'Nom principal de l’application, accompagné de la patte et du chat intégrés au visuel fourni.')
replace('Cercle vert sauge','Mascotte féline')
replace('Couleur de marque utilisée pour les repères visuels et les surfaces secondaires.',
        'Chat orange dans un médaillon bleu, repris dans les bandeaux verts des interfaces.')
replace('Nom ANIWORLD','Signature du logo')
replace('Affiché à proximité du symbole sur la couverture et dans la navigation de l’application.',
        'Les mentions THE OFFICIAL APPLICATION et HOME · DOG · CAT restent intégrées au logotype original.')
replace('Conserver les proportions du logo et son fond transparent. Prévoir un espace libre autour du symbole équivalant au quart de son diamètre. Sur l’interface, utiliser un diamètre apparent d’au moins 32 pixels ; dans un document imprimé, un diamètre d’au moins 15 mm.',
        'Conserver les proportions et la transparence du logotype. Prévoir une zone libre autour du visuel au moins égale à la hauteur de la lettre A. Dans la navigation, réserver une largeur d’environ 155 pixels ; à l’impression, privilégier une largeur d’au moins 45 mm.')
replace('Privilégier un fond blanc ou ivoire. Ne pas déformer le symbole, ne pas en modifier les couleurs et éviter les fonds chargés. Le nom de l’application reste présent lorsque le contexte ne permet pas d’identifier la marque avec le seul symbole.',
        'Privilégier un fond blanc ou ivoire. Ne pas déformer le visuel, ne pas en modifier les couleurs et éviter les fonds chargés. La mascotte peut accompagner un bandeau, tandis que le logotype complet reste visible dans la navigation.')
replace('La charte reprend la sobriété de la maquette de référence : navigation latérale, surfaces claires, cartes de synthèse, tableaux lisibles et boutons arrondis. La couleur verte marque les actions principales et les informations positives.',
        'Les six interfaces reprennent le style de la maquette de référence : barre latérale blanche avec le logo, bandeau vert avec la mascotte, cartes de statistiques, tableaux et pastilles de statut. Une colonne à droite rassemble graphiques et informations utiles. Les boutons principaux sont verts et les encarts de suivi sont jaune pâle.')
replace('La page d’accueil donne une vue immédiate de l’activité. Les indicateurs présentent les services du jour, les employés présents, les alertes de stock et les locaux disponibles.',
        'La page d’accueil donne une vue immédiate de l’activité. Les cartes présentent le total des services, les prestations planifiées, celles en cours et celles terminées. Le bandeau avec la mascotte reprend l’identité visuelle du centre.')
replace('Le tableau associe prestation, animal, horaire et statut. La sélection d’une ligne donne accès à la fiche du service.',
        'Le tableau associe prestation, horaire et statut. La sélection d’une ligne donne accès à la fiche complète du service et à l’animal concerné.')
replace('Alertes','Colonne de droite')
replace('La zone inférieure regroupe les éléments demandant une action ; chaque alerte conduit au module responsable.',
        'Un graphique présente la répartition des prestations par catégorie. Un encart jaune pâle souligne les soins terminés et les informations de suivi.')
replace('La liste affiche le nom, le poste, une compétence principale et le taux de charge. La fiche complète contient les coordonnées et les disponibilités.',
        'La liste affiche le nom, le poste et le taux de charge. La fiche complète contient les coordonnées, les compétences et les disponibilités ; un graphique compare la charge prévue.')
replace('Les colonnes montrent la désignation, le lot, la quantité et l’état. La fiche distingue les entrées, sorties et corrections.',
        'Les colonnes montrent la désignation, la quantité et l’état. La fiche détaille les lots, les entrées, les sorties et les corrections ; la colonne de droite signale le lot conseillé.')
replace('La zone de proposition indique l’horaire, la durée, le local et la ressource retenue. La confirmation réévalue les disponibilités.',
        'L’encart jaune de la colonne droite indique l’horaire, la durée, le local et la ressource retenue. La confirmation réévalue les disponibilités.')
replace('La police de l’interface est Segoe UI. Les titres d’écran utilisent 26 à 28 pixels, les titres de zone 18 pixels et le texte courant 14 à 16 pixels. Les informations secondaires restent lisibles et ne portent pas seules une consigne essentielle. La grille utilise des espacements réguliers de 8, 16 et 24 pixels.',
        'La police de l’interface est Segoe UI. Les titres d’écran utilisent 25 à 27 pixels, les titres de zone 14 à 16 pixels et le texte courant 12 à 13 pixels. Les mentions secondaires utilisent 10 à 11 pixels et ne portent pas seules une consigne essentielle. Les marges, espacements et angles arrondis restent cohérents entre les six vues.')

# Mettre à jour la palette visible sans toucher aux autres tableaux.
palette={
    '#9FC8AA':('#6D9B7B','Graphiques et accents secondaires.'),
    '#3B3B3B':('#243D3D','Texte principal et titres de l’interface.'),
    '#E6F0E7':('#DFEEE7','Bandeaux avec la mascotte et surfaces de suivi.'),
    '#9A671B':('#AC853F','Texte et icônes des encarts jaune pâle.'),
    '#A8423F':('#AE5751','Erreurs et actions nécessitant une attention.')
}
for table in doc.tables:
    if [c.text for c in table.rows[0].cells]==['Couleur','Code','Usage']:
        for row in table.rows[1:]:
            code=row.cells[1].text
            if code in palette:
                new,usage=palette[code];replace(code,new);replace(row.cells[2].text,usage)
                pr=row.cells[1]._tc.get_or_add_tcPr()
                for e in list(pr.findall(qn('w:shd'))):pr.remove(e)
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),new[1:]);pr.append(sh)
                for run in row.cells[1].paragraphs[0].runs:run.font.color.rgb=RGBColor.from_string('333333' if new=='#DFEEE7' else 'FFFFFF')

# Remplacer les illustrations dans les mêmes emplacements.
shapes=list(doc.inline_shapes)
assert len(shapes)==8
def set_size(shape,width,height):
    shape.width=width;shape.height=height
    ext=shape._inline.xpath('.//a:xfrm/a:ext')[0];ext.set('cx',str(width));ext.set('cy',str(height))

logo_image=Image.open(NEWLOGO);bounds=logo_image.getbbox();w,h=logo_image.size;l,t,r,b=bounds
for shape in shapes[:2]:
    blip=shape._inline.xpath('.//a:blip')[0];part=doc.part.related_parts[blip.get(qn('r:embed'))];part._blob=NEWLOGO.read_bytes()
    fill=blip.getparent()
    for old in list(fill.findall(qn('a:srcRect'))):fill.remove(old)
    crop=OxmlElement('a:srcRect')
    for key,value in [('l',l/w),('t',t/h),('r',(w-r)/w),('b',(h-b)/h)]:crop.set(key,str(round(value*100000)))
    fill.insert(1,crop)
    set_size(shape,Cm(8.7),Cm(8.7*(b-t)/(r-l)))
    shape._inline.docPr.set('descr','Logotype ANIWORLD fourni avec un chat orange dans un médaillon bleu')
for idx,shape in enumerate(shapes[2:]):
    path=WORK/'gui_aniworld'/f'gui_{idx}.png';im=Image.open(path);w,h=im.size
    blip=shape._inline.xpath('.//a:blip')[0];doc.part.related_parts[blip.get(qn('r:embed'))]._blob=path.read_bytes()
    set_size(shape,Cm(17),Cm(17*h/w))
    shape._inline.docPr.set('descr',f'Maquette Qt ANIWORLD numéro {idx+1}, avec le style de la référence et des données fictives')

# Le logotype complet remplace la ligne de marque redondante sur la couverture.
for p in doc.paragraphs:
    if p.text=='ANIWORLD':p._element.getparent().remove(p._element);break
for p in doc.paragraphs:
    if p._p.xpath('.//wp:inline'):
        p.paragraph_format.space_after=Pt(20);break

for part in doc.part.package.parts:
    if hasattr(part,'element'):
        for t in part.element.iter(qn('w:t')):
            if t.text:t.text=t.text.replace('de ANIWORLD','d’ANIWORLD')
doc.save(DEST)
print(DEST)
print('Updated 6 interfaces and 2 logos; retained the 5 functional module specifications.')
