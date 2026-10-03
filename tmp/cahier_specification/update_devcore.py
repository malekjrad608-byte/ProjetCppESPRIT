from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\malek\Documents\gestion_service')
WORK=ROOT/'tmp'/'cahier_specification'
SOURCE=ROOT/'output/documents/ANIWORLD_Cahier_de_specification.docx'
DEST=ROOT/'output/documents/ANIWORLD_devCore_Cahier_de_specification.docx'
LOGO=Path(r'C:\Users\malek\Downloads\logo-removebg-preview.png')
doc=Document(SOURCE)
original=list(doc.paragraphs)

def setp(p,text):
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]:r.text=''
    else:p.add_run(text)

def cell(c,text):setp(c.paragraphs[0],text)
def para(i,text):setp(original[i],text)

doc.core_properties.title='ANIWORLD — Cahier de spécification — devCore'
doc.core_properties.author='devCore'
doc.core_properties.subject='Projet C++ 2A — Gestion d’un centre de soins animaliers'
doc.core_properties.keywords='ANIWORLD, devCore, services, employés, stock, locaux, dossiers médicaux'
for part in doc.part.package.parts:
    if hasattr(part,'element'):
        for t in part.element.iter(qn('w:t')):
            if t.text:t.text=t.text.replace('Version 1.0','Version 1.1')
cover=doc.tables[0]
cell(cover.rows[2].cells[0],'Groupe')
cell(cover.rows[2].cells[1],'devCore')
cell(cover.rows[3].cells[1],'Équipe devCore')
cell(cover.rows[5].cells[1],'Version 1.1  |  27 septembre 2026')
brand=original[4].insert_paragraph_before('ANIWORLD')
brand.alignment=WD_ALIGN_PARAGRAPH.CENTER
brand.paragraph_format.space_after=Pt(7)
brand.paragraph_format.keep_with_next=True
for r in brand.runs:r.font.size=Pt(28);r.bold=True;r.font.color.rgb=RGBColor.from_string('246756')
group=original[4].insert_paragraph_before('Groupe devCore')
group.alignment=WD_ALIGN_PARAGRAPH.CENTER
group.paragraph_format.space_after=Pt(15)
group.paragraph_format.keep_with_next=True
for r in group.runs:r.font.size=Pt(13);r.bold=True;r.font.color.rgb=RGBColor.from_string('3B3B3B')

para(158,'L’identité visuelle d’ANIWORLD associe une patte anthracite, un cercle vert sauge et une interface claire. Les cinq modules partagent cette identité, avec des vues adaptées à leurs usages.')
para(161,'Figure 1  •  Logo ANIWORLD fourni pour le projet')
para(162,'Le nom ANIWORLD évoque l’univers des animaux. La patte représente les animaux pris en charge ; le cercle vert forme un repère simple et reconnaissable. Le nom de l’application accompagne ce symbole dans la navigation et sur la couverture du cahier réalisé par devCore.')
for row,values in zip(doc.tables[13].rows[1:],[
    ['Patte anthracite','Symbole principal lié aux animaux et à l’activité du centre. Sa couleur sombre assure un contraste sur le cercle vert.'],
    ['Cercle vert sauge','Couleur de marque reprise dans les repères visuels et les surfaces secondaires.'],
    ['Nom ANIWORLD','Affiché à proximité du symbole sur la couverture et dans la navigation de l’application.']]):
    for c,t in zip(row.cells,values):cell(c,t)
para(165,'Conserver les proportions du logo et son fond transparent. Prévoir une zone libre autour du symbole équivalant au quart de son diamètre. Dans l’interface, utiliser un diamètre apparent d’au moins 32 pixels ; à l’impression, un diamètre d’au moins 15 mm.')
para(166,'Privilégier un fond blanc ou ivoire. Ne pas déformer le symbole, modifier ses couleurs ou le placer sur un fond chargé. Le nom ANIWORLD reste présent lorsque le contexte ne permet pas d’identifier l’application avec le seul symbole.')
para(169,'La navigation latérale, les couleurs, la typographie et les boutons restent communs. Chaque module possède une composition adaptée : planning des services, annuaire des employés, inventaire des lots, cartes des locaux et historique médical. Le tableau de bord rassemble les indicateurs et les priorités.')
palette={'#6D9B7B':('#9FC8AA','Cercle du logo et repères de marque.'),'#243D3D':('#3B3B3B','Patte du logo et textes à fort contraste.'),'#DFEEE7':('#E3EEE5','Surfaces de synthèse et éléments sélectionnés.')}
for row in doc.tables[14].rows[1:]:
    if row.cells[1].text in palette:
        new,usage=palette[row.cells[1].text];cell(row.cells[1],new);cell(row.cells[2],usage)
        pr=row.cells[1]._tc.get_or_add_tcPr()
        for e in list(pr.findall(qn('w:shd'))):pr.remove(e)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),new[1:]);pr.append(sh)
        for r in row.cells[1].paragraphs[0].runs:r.font.color.rgb=RGBColor.from_string('FFFFFF' if new=='#3B3B3B' else '333333')

para(180,'La page d’accueil synthétise l’activité du centre. Les compteurs distinguent les services planifiés, en cours et terminés ; le total inclut aussi les annulations. Les prestations du jour, leur répartition et les priorités complètent cette vue.')
para(187,'Le planning journalier positionne les prestations par horaire et par local. La sélection d’un rendez-vous affiche sa fiche, tandis qu’une proposition aide à trouver un créneau compatible.')
para(194,'L’annuaire permet de rechercher un employé et d’ouvrir sa fiche. Les compétences, les disponibilités et la charge hebdomadaire aident le responsable à préparer les affectations.')
para(201,'L’inventaire présente les produits et leurs lots, avec les quantités, les seuils et les échéances. Deux panneaux expliquent les propositions de sortie par péremption et de réapprovisionnement. Les données sont fictives.')
para(205,'Enregistrer un mouvement ouvre la saisie d’une entrée, d’une sortie ou d’une correction. Toute sortie requiert une quantité, un motif et, lorsque pertinent, un service associé. La gestion des fiches articles reste accessible depuis l’inventaire.')
para(208,'Les cartes permettent de repérer rapidement la capacité et l’état de chaque local. La sélection d’un espace ouvre sa liste de contrôle de nettoyage ; les interventions préventives disposent d’un panneau dédié.')
para(215,'La liste des animaux reste visible à gauche du dossier sélectionné. Le vétérinaire consulte les observations datées, suit l’évolution du poids et prépare les rappels depuis une vue clinique organisée.')

gui_tables=[
    [('Navigation','La barre latérale donne accès aux cinq modules et indique la page active. Le logo et le nom ANIWORLD forment un repère commun.'),('Indicateurs','Les cartes affichent le total et la répartition par statut. Chaque compteur ouvre le périmètre correspondant, selon les droits de l’utilisateur.'),('Services du jour','Le tableau associe prestation, animal, horaire et statut. La sélection ouvre la fiche complète du service.'),('Répartition et priorités','Les barres présentent les catégories de prestations. L’encart de suivi rassemble les articles sous le seuil et les rappels vaccinaux à préparer.')],
    [('Planning des locaux','Les prestations sont positionnées sur une grille horaire. Les commandes Jour, Semaine et Liste changent la présentation ; l’export reprend la période choisie.'),('Fiche de prestation','Le panneau détaille l’animal, le responsable, le local, l’horaire, le tarif et le statut. Modifier ouvre le formulaire de la prestation sélectionnée.'),('Créneau compatible','La proposition précise une ressource, un local et une durée. La confirmation renouvelle les contrôles de disponibilité avant réservation.')],
    [('Annuaire et recherche','La recherche par nom ou poste et les filtres réduisent la liste. La sélection d’une ligne ouvre la fiche de l’employé.'),('Fiche et compétences','Le panneau présente le poste, les coordonnées, les compétences et la disponibilité. Les profils inactifs ou absents sont exclus des propositions d’affectation.'),('Charge hebdomadaire','Les barres comparent les heures planifiées aux heures disponibles. Le responsable examine une suggestion motivée avant de réaffecter une prestation.')],
    [('Inventaire des lots','Le tableau associe article, lot, quantité, seuil, échéance et état. La sélection donne accès à la traçabilité des mouvements.'),('Sortie par péremption','La proposition FEFO indique le lot utilisable dont l’échéance est la plus proche. Les lots expirés ou bloqués sont exclus ; la quantité est contrôlée à la validation.'),('Réapprovisionnement','Le panneau rend le calcul visible : consommation moyenne, délai fournisseur, stock de sécurité et quantité disponible. Le magasinier vérifie la proposition.')],
    [('Cartes des locaux','Chaque carte indique le nom, le type, la capacité et l’état du local. Le planning contrôle aussi les réservations futures sur la période demandée.'),('Liste de nettoyage','Les tâches à accomplir sont associées à un agent. La remise en service reste bloquée tant que les contrôles obligatoires ne sont pas validés.'),('Maintenance préventive','Le panneau signale les seuils d’utilisation atteints et l’intervention prévue. Une intervention confirmée réserve une plage d’indisponibilité.')],
    [('Sélection de l’animal','La recherche et la liste latérale ouvrent le dossier autorisé. L’en-tête rappelle l’identité de l’animal et les allergies déclarées.'),('Historique clinique','Les observations sont organisées par date et conservent leur auteur. Les onglets donnent accès aux consultations, vaccinations et documents.'),('Poids et échéances','La courbe restitue les pesées et leur variation. Le rappel vaccinal reste à préparer et à valider ; le signal sur le poids appelle une vérification du vétérinaire.')]
]
for table,rows in zip(doc.tables[15:21],gui_tables):
    assert len(table.rows)==len(rows)+1
    for row,values in zip(table.rows[1:],rows):
        for c,t in zip(row.cells,values):cell(c,t)

def set_size(shape,width,height):
    shape.width=width;shape.height=height
    ext=shape._inline.xpath('.//a:xfrm/a:ext')[0];ext.set('cx',str(width));ext.set('cy',str(height))
shapes=list(doc.inline_shapes)
assert len(shapes)==8
im=Image.open(LOGO);w,h=im.size;l,t,r,b=im.getbbox()
for i,shape in enumerate(shapes[:2]):
    blip=shape._inline.xpath('.//a:blip')[0]
    doc.part.related_parts[blip.get(qn('r:embed'))]._blob=LOGO.read_bytes()
    fill=blip.getparent()
    for old in list(fill.findall(qn('a:srcRect'))):fill.remove(old)
    crop=OxmlElement('a:srcRect')
    for k,v in [('l',l/w),('t',t/h),('r',(w-r)/w),('b',(h-b)/h)]:crop.set(k,str(round(v*100000)))
    fill.insert(1,crop)
    size=3.6 if i==0 else 3.7
    set_size(shape,Cm(size),Cm(size*(b-t)/(r-l)))
    shape._inline.docPr.set('descr','Logo ANIWORLD : patte anthracite sur un cercle vert sauge')
for p in doc.paragraphs:
    if p._p.xpath('.//wp:inline'):
        p.paragraph_format.space_after=Pt(10)
        break
for idx,shape in enumerate(shapes[2:]):
    path=WORK/'gui_devcore'/f'gui_{idx}.png'
    im=Image.open(path);w,h=im.size
    blip=shape._inline.xpath('.//a:blip')[0]
    doc.part.related_parts[blip.get(qn('r:embed'))]._blob=path.read_bytes()
    set_size(shape,Cm(17),Cm(17*h/w))
    shape._inline.docPr.set('descr',f'ANIWORLD — Maquette du module {idx} — Groupe devCore — Données fictives')
doc.save(DEST)
print(DEST)
